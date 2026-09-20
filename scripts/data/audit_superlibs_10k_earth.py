#!/usr/bin/env python3
import argparse
import csv
import hashlib
import json
import math
import os
import re
import shutil
import statistics
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

ENERGY_ON_TARGET = {
    "3mj": 2.4,
    "4mj": 3.2,
    "5mj": 4.0,
    "7mj": 5.6,
    "9mj": 7.2,
}
REFERENCE_SPECTRAL_COLUMNS = {
    "3mj": 330500,
    "4mj": 281625,
    "5mj": 330500,
    "7mj": 48250,
    "9mj": 48250,
}
REFERENCE_TARGETS = 2644

HEADER_ROWS = [
    "date_time",
    "carousel_number",
    "sample_id",
    "target_number",
    "location_on_sample",
    "rock_type",
    "atmosphere",
    "distance_to_target_mm",
    "locations_per_sample",
    "integrations_per_location",
    "pulses_per_integration",
    "laser_attenuation_mj",
    "laser_repetition_rate_hz",
    "integration_time_s",
    "integrations_averaged",
    "uv_trigger_delay_ns",
    "vis_trigger_delay_ns",
    "vnir_trigger_delay_ns",
    "dark_subtracted",
    "boxcar_smoothing",
    "pulses_discarded",
    "spectral_columns_discarded",
    "saturated_integrations",
]

XML_CANDIDATES = {
    "logical_identifier": ["logical_identifier"],
    "title": ["title"],
    "xml_file_name": ["file_name"],
    "xml_sample_id": ["sample_id", "sample_name", "specimen_id", "specimen_name"],
    "xml_target_name": ["target_name", "target_id"],
}


def clean(value):
    if value is None:
        return ""
    return str(value).strip().replace("\ufeff", "")


def norm(value):
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def local_name(tag):
    return tag.rsplit("}", 1)[-1]


def safe_float(value):
    value = clean(value).replace(",", "")
    if not value:
        return None
    match = re.search(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", value)
    if not match:
        return None
    try:
        return float(match.group())
    except ValueError:
        return None


def scalar(value):
    if value is None or value == "":
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return value


def write_csv(path, rows, fields=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(rows)
    if fields is None:
        fields = []
        seen = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.add(key)
                    fields.append(key)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_verification(root):
    path = root / "verification_report.txt"
    result = {"path": str(path), "exists": path.is_file(), "status": "", "full_md5": ""}
    if not path.is_file():
        return result
    text = path.read_text(encoding="utf-8", errors="replace")
    for line in text.splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key = norm(key)
        if key == "status":
            result["status"] = clean(value)
        elif key == "fullmd5verification":
            result["full_md5"] = clean(value)
        elif key == "totalselectedfiles":
            result["selected_files"] = safe_float(value)
    result["pass"] = result["status"].upper() == "PASS" and result["full_md5"].lower() == "true"
    return result


def read_manifest(root):
    paths = [
        root / "urn-nasa-pds-libs_reference_database.md5",
        root / "manifests" / "bundle_after.md5",
        root / "manifests" / "bundle_before.md5",
    ]
    path = next((x for x in paths if x.is_file()), None)
    result = {"path": str(path) if path else "", "exists": bool(path), "csv": set(), "xml": set()}
    if not path:
        return result
    prefix = "data_superlibs/10k/earth/"
    for raw in path.read_text(encoding="utf-8", errors="replace").splitlines():
        parts = raw.strip().split(maxsplit=1)
        if len(parts) != 2:
            continue
        rel = parts[1].lstrip("* ").replace("\\", "/").lstrip("/")
        if not rel.startswith(prefix):
            continue
        if rel.lower().endswith(".csv"):
            result["csv"].add(rel)
        elif rel.lower().endswith(".xml"):
            result["xml"].add(rel)
    result["csv_count"] = len(result["csv"])
    result["xml_count"] = len(result["xml"])
    return result


def read_metadata_xlsx(path):
    if load_workbook is None:
        raise RuntimeError("openpyxl is required: pip install openpyxl")
    if not path.is_file():
        raise FileNotFoundError(path)

    wb = load_workbook(path, read_only=True, data_only=True)
    candidates = []
    for ws in wb.worksheets:
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            preview.append([clean(x) for x in row])
        for i, row in enumerate(preview, 1):
            names = [norm(x) for x in row]
            if "pelletname" in names or "targetname" in names or "sampleid" in names:
                candidates.append((ws.max_row or 0, ws.max_column or 0, ws, i, row))
                break
    if not candidates:
        wb.close()
        raise RuntimeError("Could not find PELLET NAME, TARGET NAME, or SAMPLE ID in metadata workbook")

    _, _, ws, header_row, raw_headers = max(candidates, key=lambda x: (x[0], x[1]))
    headers = []
    used = Counter()
    for i, value in enumerate(raw_headers, 1):
        name = clean(value) or f"column_{i}"
        used[name] += 1
        if used[name] > 1:
            name = f"{name}_{used[name]}"
        headers.append(name)

    unit_row = header_row + 1
    units = [clean(x) for x in next(ws.iter_rows(min_row=unit_row, max_row=unit_row, values_only=True))]
    first_col = 0
    for i, value in enumerate(headers):
        if norm(value) in {"pelletname", "targetname", "sampleid"}:
            first_col = i
            break

    targets = []
    chemistry = []
    for values in ws.iter_rows(min_row=header_row + 2, values_only=True):
        values = list(values)
        if first_col >= len(values):
            continue
        target = clean(values[first_col])
        if not target:
            continue
        row = {"metadata_target_name": target, "metadata_target_key": norm(target)}
        n_values = 0
        for i, header in enumerate(headers):
            if i >= len(values) or i == first_col:
                continue
            raw = scalar(values[i])
            if raw == "":
                continue
            value = safe_float(raw)
            unit = units[i] if i < len(units) else ""
            row[header] = raw
            qualifier = ""
            text = clean(raw)
            if text.startswith("<"):
                qualifier = "<"
            elif text.startswith(">"):
                qualifier = ">"
            elif value is None:
                qualifier = "non_numeric"
            chemistry.append({
                "metadata_target_name": target,
                "metadata_target_key": norm(target),
                "analyte": header,
                "unit": unit,
                "value": "" if value is None else value,
                "qualifier": qualifier,
                "raw_value": raw,
            })
            if value is not None:
                n_values += 1
        row["numeric_metadata_values"] = n_values
        targets.append(row)

    wb.close()
    keys = defaultdict(list)
    for row in targets:
        keys[row["metadata_target_key"]].append(row["metadata_target_name"])
    collisions = {k: v for k, v in keys.items() if len(set(v)) > 1}
    unique_lookup = {k: v[0] for k, v in keys.items() if len(set(v)) == 1}
    return {
        "sheet": ws.title,
        "header_row": header_row,
        "targets": targets,
        "chemistry": chemistry,
        "lookup": unique_lookup,
        "collisions": collisions,
        "headers": headers,
        "units": units,
    }


def parse_filename(path):
    stem = path.stem
    match = re.search(r"_(t\d+l\d+)_([^/]+)_spect$", stem, flags=re.I)
    if not match:
        return {"filename_location": "", "filename_target_id": ""}
    return {
        "filename_location": match.group(1),
        "filename_target_id": match.group(2),
    }


def read_csv_header(path):
    rows = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.reader(f)
        for _ in range(24):
            try:
                rows.append(next(reader))
            except StopIteration:
                break
    if len(rows) < 24:
        raise RuntimeError(f"only {len(rows)} header rows")

    result = {}
    for i, key in enumerate(HEADER_ROWS):
        row = [clean(x) for x in rows[i]]
        first = row[0] if row else ""
        tokens = [x for x in key.split("_") if x not in {"mj", "mm", "hz", "s", "ns", "on", "to", "per"}]
        labelled = len(row) > 1 and all(norm(x) in norm(first) for x in tokens)
        values = [x for x in (row[1:] if labelled else row) if x]
        result[key] = " | ".join(values)
        result[f"header_name_{i + 1}"] = first if labelled else ""

    columns = [clean(x) for x in rows[23] if clean(x)]
    result["column_names"] = " | ".join(columns)
    result["n_columns"] = len(columns)
    result["n_spectral_columns"] = max(0, len(columns) - 1)
    return result


def read_xml(path, collect_schema=False):
    root = ET.parse(path).getroot()
    values = defaultdict(list)
    tags = Counter()
    examples = {}
    for elem in root.iter():
        if len(elem):
            continue
        tag = local_name(elem.tag)
        value = clean(elem.text)
        if not value:
            continue
        key = norm(tag)
        values[key].append(value)
        if collect_schema:
            tags[tag] += 1
            examples.setdefault(tag, value[:200])

    result = {}
    for out_key, candidates in XML_CANDIDATES.items():
        result[out_key] = ""
        for candidate in candidates:
            found = values.get(norm(candidate), [])
            if found:
                result[out_key] = found[0]
                break
    return result, tags, examples


def grid_signature(path):
    h = hashlib.sha256()
    count = 0
    first = None
    last = None
    minimum = None
    maximum = None
    nonnumeric = 0
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.reader(f)
        for _ in range(24):
            next(reader, None)
        for row in reader:
            if not row:
                continue
            value = safe_float(row[0])
            if value is None:
                nonnumeric += 1
                continue
            token = f"{value:.12g}\n".encode("ascii")
            h.update(token)
            count += 1
            if first is None:
                first = value
            last = value
            minimum = value if minimum is None else min(minimum, value)
            maximum = value if maximum is None else max(maximum, value)
    return {
        "grid_points": count,
        "wavelength_first_nm": first,
        "wavelength_last_nm": last,
        "wavelength_min_nm": minimum,
        "wavelength_max_nm": maximum,
        "wavelength_sha256": h.hexdigest(),
        "nonnumeric_data_rows": nonnumeric,
    }


def choose_grid_files(paths, mode, per_energy):
    if mode == "none":
        return set()
    if mode == "all":
        return set(paths)
    grouped = defaultdict(list)
    for path in paths:
        grouped[path.parent.name.lower()].append(path)
    selected = set()
    for energy, items in grouped.items():
        items = sorted(items)
        n = min(per_energy, len(items))
        if n == 0:
            continue
        if n == 1:
            selected.add(items[0])
            continue
        for i in range(n):
            index = round(i * (len(items) - 1) / (n - 1))
            selected.add(items[index])
    return selected


def schema_files(paths, per_energy):
    return choose_grid_files(paths, "sample", per_energy)


def match_target(candidates, lookup, collisions):
    seen = set()
    for source, value in candidates:
        key = norm(value)
        if not key or key in seen:
            continue
        seen.add(key)
        if key in collisions:
            return "", key, "ambiguous", source
        if key in lookup:
            return lookup[key], key, "matched", source
    return "", "", "unmatched", ""


def process_product(args):
    path, data_root, grid_paths, schema_paths, lookup, collisions = args
    rel = path.relative_to(data_root.parent.parent.parent)
    energy = path.parent.name.lower()
    xml_path = path.with_suffix(".xml")
    row = {
        "csv_path": str(rel),
        "xml_path": str(xml_path.relative_to(data_root.parent.parent.parent)) if xml_path.exists() else "",
        "nominal_energy_folder": energy,
        "nominal_attenuation_mj": safe_float(energy),
        "actual_energy_on_target_mj": ENERGY_ON_TARGET.get(energy, ""),
        "csv_bytes": path.stat().st_size,
        "xml_exists": xml_path.is_file(),
        "csv_parse_ok": False,
        "xml_parse_ok": False,
        "grid_checked": path in grid_paths,
    }
    row.update(parse_filename(path))
    errors = []
    tags = Counter()
    examples = {}

    try:
        row.update(read_csv_header(path))
        row["csv_parse_ok"] = True
    except Exception as e:
        errors.append(f"CSV: {e}")

    if xml_path.is_file():
        try:
            xml_values, tags, examples = read_xml(xml_path, path in schema_paths)
            row.update(xml_values)
            row["xml_parse_ok"] = True
        except Exception as e:
            errors.append(f"XML: {e}")

    if path in grid_paths:
        try:
            row.update(grid_signature(path))
        except Exception as e:
            errors.append(f"GRID: {e}")

    candidates = [
        ("csv_sample_id", row.get("sample_id", "")),
        ("xml_sample_id", row.get("xml_sample_id", "")),
        ("xml_target_name", row.get("xml_target_name", "")),
        ("filename_target_id", row.get("filename_target_id", "")),
    ]
    matched, key, status, source = match_target(candidates, lookup, collisions)
    row["metadata_target_name"] = matched
    row["metadata_target_key"] = key
    row["metadata_match_status"] = status
    row["metadata_match_source"] = source
    row["error"] = " ; ".join(errors)
    return row, tags, examples


def summarize_targets(products, chemistry_counts):
    groups = defaultdict(list)
    for row in products:
        key = row.get("metadata_target_name") or row.get("sample_id") or row.get("filename_target_id")
        key = clean(key)
        if not key:
            key = "__UNKNOWN__"
        groups[key].append(row)

    rows = []
    for target, items in sorted(groups.items(), key=lambda x: x[0].lower()):
        energies = sorted({clean(x.get("nominal_energy_folder")) for x in items if clean(x.get("nominal_energy_folder"))})
        locations = sorted({clean(x.get("location_on_sample")) or clean(x.get("filename_location")) for x in items if clean(x.get("location_on_sample")) or clean(x.get("filename_location"))})
        rock_types = sorted({clean(x.get("rock_type")) for x in items if clean(x.get("rock_type"))})
        matched = [x for x in items if x.get("metadata_match_status") == "matched"]
        row = {
            "target_group": target,
            "metadata_target_name": matched[0].get("metadata_target_name", "") if matched else "",
            "metadata_match_status": "matched" if matched else items[0].get("metadata_match_status", "unmatched"),
            "n_products": len(items),
            "n_spectral_columns": sum(int(x.get("n_spectral_columns") or 0) for x in items),
            "n_energies": len(energies),
            "energies": " | ".join(energies),
            "n_locations": len(locations),
            "locations": " | ".join(locations),
            "rock_types": " | ".join(rock_types),
            "n_chemistry_values": chemistry_counts.get(norm(target), 0),
            "csv_parse_failures": sum(not bool(x.get("csv_parse_ok")) for x in items),
            "xml_parse_failures": sum(not bool(x.get("xml_parse_ok")) for x in items),
        }
        for energy in sorted(ENERGY_ON_TARGET):
            row[f"products_{energy}"] = sum(x.get("nominal_energy_folder") == energy for x in items)
            row[f"spectra_{energy}"] = sum(int(x.get("n_spectral_columns") or 0) for x in items if x.get("nominal_energy_folder") == energy)
        rows.append(row)
    return rows


def analyte_coverage(chemistry, matched_names):
    matched_keys = {norm(x) for x in matched_names}
    grouped = defaultdict(list)
    for row in chemistry:
        if row["metadata_target_key"] in matched_keys and row["value"] != "":
            grouped[(row["analyte"], row["unit"])].append(row)
    out = []
    denominator = len(matched_keys)
    for (analyte, unit), items in sorted(grouped.items()):
        values = [x["value"] for x in items]
        targets = {x["metadata_target_key"] for x in items}
        out.append({
            "analyte": analyte,
            "unit": unit,
            "matched_targets_with_value": len(targets),
            "matched_targets_total": denominator,
            "coverage_fraction": len(targets) / denominator if denominator else 0,
            "min": min(values),
            "median": statistics.median(values),
            "max": max(values),
            "qualified_values": sum(bool(x["qualifier"]) for x in items),
        })
    return out


def make_report(path, summary, top_analytes, warnings):
    lines = [
        "# SuperLIBS 10K Earth audit report",
        "",
        f"- Root: `{summary['root']}`",
        f"- Verification report accepted: **{summary['verification_pass']}**",
        f"- CSV products found: **{summary['csv_products']:,}**",
        f"- XML labels paired: **{summary['xml_pairs']:,}**",
        f"- Unique product target groups: **{summary['target_groups']:,}**",
        f"- Unique metadata targets matched: **{summary['matched_metadata_targets']:,}**",
        f"- CSV parse failures: **{summary['csv_parse_failures']:,}**",
        f"- XML parse failures: **{summary['xml_parse_failures']:,}**",
        f"- Missing XML labels: **{summary['missing_xml']:,}**",
        f"- Wavelength grids checked: **{summary['grids_checked']:,}**",
        f"- Distinct checked wavelength hashes: **{summary['distinct_grid_hashes']:,}**",
        "",
        "## Energy coverage",
        "",
        "| Folder | Nominal setting (mJ) | Energy on target (mJ) | Products | Spectral columns | PDS guide count | Delta |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for energy, values in summary["energy_coverage"].items():
        lines.append(
            f"| {energy} | {values['nominal_attenuation_mj']} | {values['actual_energy_on_target_mj']} | "
            f"{values['products']:,} | {values['spectral_columns']:,} | {values['reference_spectral_columns']:,} | "
            f"{values['spectral_column_delta']:,} |"
        )

    lines += ["", "## Highest chemistry coverage among matched targets", "", "| Analyte | Unit | Targets | Coverage |", "|---|---|---:|---:|"]
    for row in top_analytes[:20]:
        lines.append(f"| {row['analyte']} | {row['unit']} | {row['matched_targets_with_value']:,} | {row['coverage_fraction']:.1%} |")

    lines += ["", "## Warnings and decisions required", ""]
    if warnings:
        lines.extend(f"- {x}" for x in warnings)
    else:
        lines.append("- No structural blockers were detected by this audit.")

    lines += [
        "",
        "## Output files",
        "",
        "- `audit_summary.json`: machine-readable totals and readiness flags.",
        "- `products.csv`: one row per spectral CSV/XML product pair.",
        "- `targets.csv`: target-level grouping and energy/location coverage.",
        "- `metadata_targets.csv`: rows parsed from `libs_metadata.xlsx`.",
        "- `chemistry_long.csv`: analyte values with original units and qualifiers.",
        "- `analyte_coverage.csv`: chemistry coverage among matched SuperLIBS targets.",
        "- `wavelength_grids.csv`: wavelength signatures for checked products.",
        "- `xml_fields.csv`: XML leaf fields observed in schema samples.",
        "- `errors.csv`: products with missing pairs or parser/grid errors.",
        "- `unmatched_targets.csv`: target groups not mapped uniquely to the metadata workbook.",
        "- `manifest_differences.csv`: missing or extra CSV/XML paths relative to the frozen PDS manifest.",
        "",
        "Do not build the definitive HDF5 until unmatched or ambiguous target identities and any wavelength-grid differences have been reviewed.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="Audit and index the verified NASA SuperLIBS 10K Earth collection.")
    parser.add_argument("root", nargs="?", default="/home/user/roboai_green/nasa_calibration")
    parser.add_argument("--out", default="")
    parser.add_argument("--workers", type=int, default=max(1, min(16, os.cpu_count() or 4)))
    parser.add_argument("--grid-check", choices=["none", "sample", "all"], default="sample")
    parser.add_argument("--grid-samples-per-energy", type=int, default=20)
    parser.add_argument("--xml-schema-samples-per-energy", type=int, default=5)
    parser.add_argument("--max-products", type=int, default=0, help="Development/smoke-test limit; 0 scans all products.")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    root = Path(args.root).expanduser().resolve()
    out = Path(args.out).expanduser().resolve() if args.out else root / "audit_superlibs_10k_earth"
    data_root = root / "data_superlibs" / "10k" / "earth"
    metadata_path = root / "document" / "libs_metadata.xlsx"

    if not data_root.is_dir():
        raise SystemExit(f"ERROR: data directory not found: {data_root}")
    if out.exists() and any(out.iterdir()):
        if not args.overwrite:
            raise SystemExit(f"ERROR: output directory is not empty: {out}\nUse --overwrite to replace audit outputs.")
        shutil.rmtree(out)
    out.mkdir(parents=True, exist_ok=True)

    verification = read_verification(root)
    manifest = read_manifest(root)
    metadata = read_metadata_xlsx(metadata_path)

    csv_paths = sorted(data_root.glob("*/*.csv"))
    if args.max_products:
        csv_paths = csv_paths[:args.max_products]
    if not csv_paths:
        raise SystemExit(f"ERROR: no CSV products found under {data_root}")

    grid_paths = choose_grid_files(csv_paths, args.grid_check, args.grid_samples_per_energy)
    schema_paths = schema_files(csv_paths, args.xml_schema_samples_per_energy)

    print(f"Root: {root}")
    print(f"Output: {out}")
    print(f"Products to scan: {len(csv_paths):,}")
    print(f"Grid check: {args.grid_check} ({len(grid_paths):,} products)")
    print(f"Workers: {args.workers}")

    work = ((p, data_root, grid_paths, schema_paths, metadata["lookup"], metadata["collisions"]) for p in csv_paths)
    products = []
    xml_fields = Counter()
    xml_examples = {}
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        for i, (row, tags, examples) in enumerate(pool.map(process_product, work), 1):
            products.append(row)
            xml_fields.update(tags)
            for key, value in examples.items():
                xml_examples.setdefault(key, value)
            if i % 500 == 0 or i == len(csv_paths):
                print(f"  scanned {i:,}/{len(csv_paths):,}", flush=True)

    chemistry_counts = Counter(x["metadata_target_key"] for x in metadata["chemistry"] if x["value"] != "")
    targets = summarize_targets(products, chemistry_counts)
    matched_names = sorted({x["metadata_target_name"] for x in products if x.get("metadata_match_status") == "matched"})
    coverage = analyte_coverage(metadata["chemistry"], matched_names)

    grid_rows = []
    for row in products:
        if row.get("grid_checked"):
            grid_rows.append({
                "csv_path": row["csv_path"],
                "nominal_energy_folder": row["nominal_energy_folder"],
                "grid_points": row.get("grid_points", ""),
                "wavelength_first_nm": row.get("wavelength_first_nm", ""),
                "wavelength_last_nm": row.get("wavelength_last_nm", ""),
                "wavelength_min_nm": row.get("wavelength_min_nm", ""),
                "wavelength_max_nm": row.get("wavelength_max_nm", ""),
                "wavelength_sha256": row.get("wavelength_sha256", ""),
                "nonnumeric_data_rows": row.get("nonnumeric_data_rows", ""),
                "error": row.get("error", ""),
            })

    errors = []
    for row in products:
        if row.get("error") or not row.get("xml_exists"):
            errors.append({
                "csv_path": row["csv_path"],
                "xml_path": row.get("xml_path", ""),
                "xml_exists": row.get("xml_exists", False),
                "csv_parse_ok": row.get("csv_parse_ok", False),
                "xml_parse_ok": row.get("xml_parse_ok", False),
                "metadata_match_status": row.get("metadata_match_status", ""),
                "error": row.get("error", "") or "Missing XML label",
            })

    energy_coverage = {}
    for energy in sorted({x["nominal_energy_folder"] for x in products}):
        items = [x for x in products if x["nominal_energy_folder"] == energy]
        spectral_columns = sum(int(x.get("n_spectral_columns") or 0) for x in items)
        reference_columns = REFERENCE_SPECTRAL_COLUMNS.get(energy, "")
        energy_coverage[energy] = {
            "nominal_attenuation_mj": safe_float(energy),
            "actual_energy_on_target_mj": ENERGY_ON_TARGET.get(energy, ""),
            "products": len(items),
            "spectral_columns": spectral_columns,
            "reference_spectral_columns": reference_columns,
            "spectral_column_delta": spectral_columns - reference_columns if reference_columns != "" else "",
            "matched_products": sum(x.get("metadata_match_status") == "matched" for x in items),
        }

    grid_hashes = {x.get("wavelength_sha256") for x in grid_rows if x.get("wavelength_sha256")}
    complete_scan = not args.max_products
    actual_csv = {x["csv_path"] for x in products}
    actual_xml = {x["xml_path"] for x in products if x.get("xml_path")}
    manifest_missing_csv = sorted(manifest["csv"] - actual_csv) if manifest.get("exists") and complete_scan else []
    manifest_extra_csv = sorted(actual_csv - manifest["csv"]) if manifest.get("exists") and complete_scan else []
    manifest_missing_xml = sorted(manifest["xml"] - actual_xml) if manifest.get("exists") and complete_scan else []
    manifest_extra_xml = sorted(actual_xml - manifest["xml"]) if manifest.get("exists") and complete_scan else []
    manifest_match = bool(manifest.get("exists")) and complete_scan and not any([
        manifest_missing_csv, manifest_extra_csv, manifest_missing_xml, manifest_extra_xml
    ])
    part_files = list(root.rglob("*.part"))
    warnings = []
    missing_xml = sum(not x.get("xml_exists") for x in products)
    csv_fail = sum(not x.get("csv_parse_ok") for x in products)
    xml_fail = sum(not x.get("xml_parse_ok") for x in products)
    unmatched_groups = sum(x.get("metadata_match_status") != "matched" for x in targets)
    if not verification.get("pass"):
        warnings.append("The local verification report does not state both `Status: PASS` and `Full MD5 verification: True`.")
    if not manifest.get("exists"):
        warnings.append("The frozen PDS MD5 manifest was not found in the dataset root or `manifests/` directory.")
    elif complete_scan and not manifest_match:
        warnings.append("The scanned CSV/XML inventory differs from the frozen PDS manifest; inspect `manifest_differences.csv`.")
    if part_files:
        warnings.append(f"{len(part_files):,} `.part` files remain under the dataset root.")
    if missing_xml:
        warnings.append(f"{missing_xml:,} CSV products have no paired XML label.")
    if csv_fail:
        warnings.append(f"{csv_fail:,} CSV headers could not be parsed.")
    if xml_fail:
        warnings.append(f"{xml_fail:,} XML labels could not be parsed.")
    if unmatched_groups:
        warnings.append(f"{unmatched_groups:,} target groups did not map uniquely to the metadata workbook; inspect `targets.csv`.")
    if metadata["collisions"]:
        warnings.append(f"{len(metadata['collisions']):,} normalized metadata target keys are ambiguous.")
    if len(grid_hashes) > 1:
        warnings.append(f"The checked products contain {len(grid_hashes):,} distinct wavelength grids; HDF5 construction must preserve or reconcile them explicitly.")
    if complete_scan and len(targets) != REFERENCE_TARGETS:
        warnings.append(f"The parser found {len(targets):,} target groups; the PDS user guide reports {REFERENCE_TARGETS:,} SuperLIBS 10K Earth reference targets.")
    for energy, values in energy_coverage.items():
        if complete_scan and values["spectral_column_delta"] not in {"", 0}:
            warnings.append(
                f"{energy} contains {values['spectral_columns']:,} spectral columns; "
                f"the PDS user guide reports {values['reference_spectral_columns']:,}."
            )
    if args.grid_check == "sample":
        warnings.append("Wavelength checking used a deterministic sample. Run again with `--grid-check all` before freezing the final HDF5 parser.")
    if args.max_products:
        warnings.append(f"This was a limited scan of {len(csv_paths):,} products because `--max-products` was used.")

    summary = {
        "root": str(root),
        "output": str(out),
        "verification": verification,
        "verification_pass": bool(verification.get("pass")),
        "manifest": {
            "path": manifest.get("path", ""),
            "exists": bool(manifest.get("exists")),
            "expected_csv": manifest.get("csv_count", 0),
            "expected_xml": manifest.get("xml_count", 0),
            "inventory_match": manifest_match,
            "missing_csv": len(manifest_missing_csv),
            "extra_csv": len(manifest_extra_csv),
            "missing_xml": len(manifest_missing_xml),
            "extra_xml": len(manifest_extra_xml),
        },
        "part_files": len(part_files),
        "complete_scan": complete_scan,
        "metadata_workbook": str(metadata_path),
        "metadata_sheet": metadata["sheet"],
        "metadata_targets": len(metadata["targets"]),
        "metadata_normalization_collisions": len(metadata["collisions"]),
        "csv_products": len(products),
        "xml_pairs": sum(x.get("xml_exists") for x in products),
        "missing_xml": missing_xml,
        "csv_parse_failures": csv_fail,
        "xml_parse_failures": xml_fail,
        "target_groups": len(targets),
        "reference_target_count": REFERENCE_TARGETS,
        "target_count_delta": len(targets) - REFERENCE_TARGETS if complete_scan else "",
        "matched_metadata_targets": len(matched_names),
        "unmatched_target_groups": unmatched_groups,
        "grids_checked": len(grid_rows),
        "distinct_grid_hashes": len(grid_hashes),
        "grid_check_mode": args.grid_check,
        "energy_coverage": energy_coverage,
        "structural_ready_for_hdf5_design": bool(verification.get("pass")) and manifest_match and not part_files and not missing_xml and not csv_fail and not xml_fail,
        "warnings": warnings,
    }

    product_fields = [
        "csv_path", "xml_path", "nominal_energy_folder", "nominal_attenuation_mj", "actual_energy_on_target_mj",
        "csv_bytes", "xml_exists", "csv_parse_ok", "xml_parse_ok", "sample_id", "target_number",
        "location_on_sample", "filename_location", "filename_target_id", "rock_type", "atmosphere",
        "distance_to_target_mm", "locations_per_sample", "integrations_per_location", "pulses_per_integration",
        "laser_attenuation_mj", "laser_repetition_rate_hz", "integration_time_s", "integrations_averaged",
        "uv_trigger_delay_ns", "vis_trigger_delay_ns", "vnir_trigger_delay_ns", "dark_subtracted",
        "boxcar_smoothing", "pulses_discarded", "spectral_columns_discarded", "saturated_integrations",
        "n_columns", "n_spectral_columns", "logical_identifier", "title", "xml_sample_id", "xml_target_name",
        "metadata_target_name", "metadata_target_key", "metadata_match_status", "metadata_match_source",
        "grid_checked", "grid_points", "wavelength_first_nm", "wavelength_last_nm", "wavelength_sha256", "error",
    ]
    write_csv(out / "products.csv", products, product_fields)
    write_csv(out / "targets.csv", targets)
    write_csv(out / "metadata_targets.csv", metadata["targets"])
    write_csv(out / "chemistry_long.csv", metadata["chemistry"], ["metadata_target_name", "metadata_target_key", "analyte", "unit", "value", "qualifier", "raw_value"])
    write_csv(out / "analyte_coverage.csv", coverage)
    write_csv(out / "wavelength_grids.csv", grid_rows)
    write_csv(out / "errors.csv", errors)
    write_csv(out / "unmatched_targets.csv", [x for x in targets if x.get("metadata_match_status") != "matched"])
    manifest_rows = []
    for status, paths in [
        ("missing_csv", manifest_missing_csv), ("extra_csv", manifest_extra_csv),
        ("missing_xml", manifest_missing_xml), ("extra_xml", manifest_extra_xml),
    ]:
        manifest_rows.extend({"status": status, "path": x} for x in paths)
    write_csv(out / "manifest_differences.csv", manifest_rows, ["status", "path"])
    write_csv(out / "xml_fields.csv", [
        {"xml_leaf_field": key, "observed_count_in_schema_samples": count, "example_value": xml_examples.get(key, "")}
        for key, count in sorted(xml_fields.items())
    ])
    (out / "metadata_key_collisions.json").write_text(json.dumps(metadata["collisions"], indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "audit_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    top = sorted(coverage, key=lambda x: (-x["matched_targets_with_value"], x["analyte"]))
    make_report(out / "audit_report.md", summary, top, warnings)

    print("\nAudit complete.")
    print(f"Report: {out / 'audit_report.md'}")
    print(f"Summary: {out / 'audit_summary.json'}")
    print(f"Structural ready for HDF5 design: {summary['structural_ready_for_hdf5_design']}")
    if warnings:
        print(f"Warnings: {len(warnings)}")


if __name__ == "__main__":
    main()
